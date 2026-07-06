"""
hist_checklist.py  -  Historial de Listas de Chequeo
=====================================================
Ventana que muestra todas las listas de chequeo guardadas,
con filtros por código de producto y lote, y exportación a Excel.
"""

import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from config.config import (
    CHECKLISTS_FILE,
    COLORS,
    PROVEEDORES_FILE,
    REPORTES_PATH,
    SUSTANCIAS_FILE,
)
from ui.styles import build_header
from ui.window_utils import maximize_window
from utils.data_handler import DataHandler, Lookups, build_substance_indexes

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    _OPENPYXL = True
except ImportError:
    _OPENPYXL = False

CHECK_ITEMS_ORDER = [
    "Nombre",
    "No. de Lote",
    "Cantidad",
    "Rótulo de Identificación",
    "Fecha de Fabricación",
    "Fecha de Vencimiento",
    "Fabricante",
    "Rótulos de seguridad, sellos y precintos de seguridad y garantía",
    "Ficha de Seguridad",
    "Certificado de Calidad",
    "Se evidencian Golpes, Roturas u Otros",
    "Cumple con las especificaciones requeridas",
]


class HistCheckListWindow:
    """Vista de historial de listas de chequeo con filtros y exportación a Excel."""

    def __init__(self, parent: tk.Tk):
        self.window = tk.Toplevel(parent)
        self.window.title("Historial de Listas de Chequeo")
        self.window.geometry("1440x600")
        self.window.configure(bg=COLORS["secondary"])
        maximize_window(self.window)

        self.filter_codigo_var = tk.StringVar()
        self.filter_lote_var = tk.StringVar()
        self.filter_estado_var = tk.StringVar(value="Todos")

        self.tree: ttk.Treeview | None = None
        self._all_rows: list[dict] = []
        self._codigo_options: list[str] = []

        # Cargar catálogos
        proveedores = DataHandler.load_json(PROVEEDORES_FILE).get("maestrasProveedores", [])
        sustancias = DataHandler.load_json(SUSTANCIAS_FILE).get("maestrasSustancias", [])
        self.lkp = Lookups(proveedores=proveedores)
        self.sustancias_by_id, self.sustancias_by_code = build_substance_indexes(sustancias)
        self._codigo_options = sorted([
            str(s.get("codigo", "")).strip()
            for s in sustancias
            if str(s.get("codigo", "")).strip() and bool(s.get("habilitada", True))
        ])

        self._build_ui()
        self.window.bind("<Escape>", lambda _e: self.window.destroy())
        self.load_table()

    # ------------------------------------------------------------------
    # UI
    # ------------------------------------------------------------------

    def _build_ui(self) -> None:
        wrapper = tk.Frame(self.window, bg="white", bd=1, relief="solid", padx=12, pady=12)
        wrapper.pack(expand=True, fill="both", padx=14, pady=14)

        build_header(wrapper, "Sistema de Gestión  -  Historial de Listas de Chequeo")

        # --- Filtros ---
        filter_row = tk.Frame(wrapper, bg="white")
        filter_row.pack(fill="x", pady=(0, 8))

        tk.Label(filter_row, text="Código:", bg="white", font=("Segoe UI", 10)).pack(side="left")
        codigo_combo = ttk.Combobox(
            filter_row,
            textvariable=self.filter_codigo_var,
            values=[""] + self._codigo_options,
            state="normal",
            width=18,
        )
        codigo_combo.pack(side="left", padx=(4, 12))
        codigo_combo.bind("<<ComboboxSelected>>", lambda _e: self.load_table())
        self.filter_codigo_var.trace_add("write", lambda *_: self.load_table())

        tk.Label(filter_row, text="Lote:", bg="white", font=("Segoe UI", 10)).pack(side="left")
        lote_entry = tk.Entry(filter_row, textvariable=self.filter_lote_var, width=16)
        lote_entry.pack(side="left", padx=(4, 12))
        lote_entry.bind("<Return>", lambda _e: self.load_table())
        self.filter_lote_var.trace_add("write", lambda *_: self.load_table())

        tk.Label(filter_row, text="Estado:", bg="white", font=("Segoe UI", 10)).pack(side="left")
        ttk.Combobox(
            filter_row,
            textvariable=self.filter_estado_var,
            values=["Todos", "ACTIVO"],
            state="readonly",
            width=10,
        ).pack(side="left", padx=(4, 12))
        self.filter_estado_var.trace_add("write", lambda *_: self.load_table())

        tk.Button(
            filter_row,
            text="⟳ Limpiar filtros",
            command=self._clear_filters,
            bg=COLORS["border"],
            fg=COLORS["text_dark"],
            relief="flat",
            padx=12,
            pady=4,
        ).pack(side="left", padx=(0, 8))

        tk.Button(
            filter_row,
            text="📥 Exportar a Excel",
            command=self._export_excel,
            bg=COLORS["success"],
            fg="white",
            relief="flat",
            font=("Segoe UI", 10, "bold"),
            padx=14,
            pady=4,
        ).pack(side="right")

        # --- Tabla ---
        columns = (
            "id", "fecha", "proveedor", "orden_compra",
            "codigo", "nombre", "lote", "cantidad",
            "aprobo", "verifico", "usuario", "estado",
        )
        headings = {
            "id": "ID",
            "fecha": "Fecha Recepción",
            "proveedor": "Proveedor",
            "orden_compra": "Orden Compra",
            "codigo": "Código",
            "nombre": "Nombre Producto",
            "lote": "Lote",
            "cantidad": "Cantidad",
            "aprobo": "Aprobó",
            "verifico": "Verificó",
            "usuario": "Usuario",
            "estado": "Estado",
        }
        widths = {
            "id": 50,
            "fecha": 110,
            "proveedor": 160,
            "orden_compra": 120,
            "codigo": 90,
            "nombre": 200,
            "lote": 100,
            "cantidad": 80,
            "aprobo": 120,
            "verifico": 120,
            "usuario": 100,
            "estado": 70,
        }

        tree_frame = tk.Frame(wrapper, bg="white")
        tree_frame.pack(expand=True, fill="both")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)

        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=18)
        style = ttk.Style(self.window)
        style.configure("Treeview", background="white", foreground="black", rowheight=24, fieldbackground="white")
        style.map("Treeview", background=[("selected", COLORS["primary"])], foreground=[("selected", "white")])

        sy = ttk.Scrollbar(tree_frame, orient="vertical", command=self.tree.yview)
        sx = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        sy.grid(row=0, column=1, sticky="ns")
        sx.grid(row=1, column=0, sticky="ew")

        for col in columns:
            self.tree.heading(col, text=headings[col], command=lambda c=col: self._sort_column(c, False))
            self.tree.column(col, width=widths[col], anchor="w", minwidth=50, stretch=True)

        self.tree.tag_configure("par", background="#FAFAFA")
        self.tree.tag_configure("impar", background="white")

        # --- Contador + Salir ---
        bottom = tk.Frame(wrapper, bg="white")
        bottom.pack(fill="x", pady=(6, 0))
        self._count_label = tk.Label(bottom, text="", bg="white", fg="#666", font=("Segoe UI", 9))
        self._count_label.pack(side="left")
        tk.Button(
            bottom,
            text="Salir",
            command=self.window.destroy,
            bg=COLORS["primary"],
            fg=COLORS["text_light"],
            relief="flat",
            padx=24,
            pady=5,
        ).pack(side="right")

    # ------------------------------------------------------------------
    # Carga de datos
    # ------------------------------------------------------------------

    def _resolve_nombre(self, record: dict) -> str:
        """Resuelve el nombre del producto desde id_sustancia o codigo_producto."""
        sid = record.get("id_sustancia")
        if sid is not None:
            sust = self.sustancias_by_id.get(sid)
            if sust:
                return str(sust.get("nombre", ""))
        codigo = str(record.get("codigo_producto", "")).strip()
        if codigo:
            from utils.data_handler import _norm
            sust = self.sustancias_by_code.get(_norm(codigo))
            if sust:
                return str(sust.get("nombre", ""))
        return ""

    def load_table(self) -> None:
        if self.tree is None:
            return

        checklists = DataHandler.get_all(CHECKLISTS_FILE, "listasChequeoRecepcionCompra")
        self._all_rows = checklists

        codigo_filter = self.filter_codigo_var.get().strip().lower()
        lote_filter = self.filter_lote_var.get().strip().lower()
        estado_filter = self.filter_estado_var.get()

        self.tree.delete(*self.tree.get_children())
        count = 0
        for idx, record in enumerate(checklists):
            codigo = str(record.get("codigo_producto", "")).strip()
            lote = str(record.get("lote", "")).strip()
            estado = str(record.get("estado", "ACTIVO")).strip()

            if codigo_filter and codigo_filter not in codigo.lower():
                continue
            if lote_filter and lote_filter not in lote.lower():
                continue
            if estado_filter != "Todos" and estado != estado_filter:
                continue

            proveedor_id = record.get("id_proveedor")
            proveedor_nombre = self.lkp.to_name("proveedores", proveedor_id) or str(record.get("proveedor", ""))
            nombre = self._resolve_nombre(record)

            row = (
                record.get("id", ""),
                str(record.get("fecha_recepcion", "")),
                proveedor_nombre,
                str(record.get("orden_compra", "")),
                codigo,
                nombre,
                lote,
                str(record.get("cantidad", "")),
                str(record.get("aprobo", "")),
                str(record.get("verifico", "")),
                str(record.get("usuario", "")),
                estado,
            )
            tag = "par" if idx % 2 == 0 else "impar"
            self.tree.insert("", tk.END, values=row, tags=(tag,))
            count += 1

        self._count_label.config(text=f"{count} registro(s) encontrado(s)")

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------

    def _clear_filters(self) -> None:
        self.filter_codigo_var.set("")
        self.filter_lote_var.set("")
        self.filter_estado_var.set("Todos")

    def _sort_column(self, col: str, reverse: bool) -> None:
        if self.tree is None:
            return
        col_idx = self.tree["columns"].index(col)
        items = [(self.tree.set(iid, col), iid) for iid in self.tree.get_children("")]
        items.sort(key=lambda x: x[0].lower(), reverse=reverse)
        for i, (_val, iid) in enumerate(items):
            self.tree.move(iid, "", i)
        self.tree.heading(col, command=lambda: self._sort_column(col, not reverse))

    # ------------------------------------------------------------------
    # Exportación a Excel
    # ------------------------------------------------------------------

    def _get_filtered_records(self) -> list[dict]:
        """Retorna los registros actualmente filtrados."""
        checklists = self._all_rows
        codigo_filter = self.filter_codigo_var.get().strip().lower()
        lote_filter = self.filter_lote_var.get().strip().lower()
        estado_filter = self.filter_estado_var.get()

        result = []
        for record in checklists:
            codigo = str(record.get("codigo_producto", "")).strip()
            lote = str(record.get("lote", "")).strip()
            estado = str(record.get("estado", "ACTIVO")).strip()

            if codigo_filter and codigo_filter not in codigo.lower():
                continue
            if lote_filter and lote_filter not in lote.lower():
                continue
            if estado_filter != "Todos" and estado != estado_filter:
                continue
            result.append(record)
        return result

    def _export_excel(self) -> None:
        if not _OPENPYXL:
            messagebox.showerror(
                "Error",
                "openpyxl no está instalado.\nInstala con: pip install openpyxl",
                parent=self.window,
            )
            return

        records = self._get_filtered_records()
        if not records:
            messagebox.showwarning("Sin datos", "No hay registros para exportar con los filtros actuales.", parent=self.window)
            return

        # Ruta de destino
        REPORTES_PATH.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"HistCheckList_{timestamp}.xlsx"
        filepath = filedialog.asksaveasfilename(
            parent=self.window,
            title="Guardar reporte de checklists",
            initialdir=str(REPORTES_PATH),
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not filepath:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Historial CheckList"

        # -- Estilos --
        header_font = Font(bold=True, color="FFFFFF", size=10)
        header_fill = PatternFill("solid", fgColor="C94A7F")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        center_align = Alignment(horizontal="center", vertical="center")
        si_fill = PatternFill("solid", fgColor="DFF3E3")
        no_fill = PatternFill("solid", fgColor="FADBD8")
        none_fill = PatternFill("solid", fgColor="FFF9C4")

        # -- Encabezados principales --
        main_cols = [
            "ID", "Fecha Recepción", "Proveedor", "Orden Compra",
            "Código", "Nombre Producto", "Lote", "Cantidad",
            "Aprobó", "Verificó", "Usuario", "Estado",
            "Observaciones",
        ]
        check_cols = CHECK_ITEMS_ORDER
        all_headers = main_cols + check_cols

        for col_idx, header in enumerate(all_headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # -- Datos --
        for row_idx, record in enumerate(records, start=2):
            proveedor_id = record.get("id_proveedor")
            proveedor_nombre = self.lkp.to_name("proveedores", proveedor_id) or str(record.get("proveedor", ""))
            nombre = self._resolve_nombre(record)
            checklist_items: dict = record.get("checklist", {})

            main_values = [
                record.get("id", ""),
                str(record.get("fecha_recepcion", "")),
                proveedor_nombre,
                str(record.get("orden_compra", "")),
                str(record.get("codigo_producto", "")),
                nombre,
                str(record.get("lote", "")),
                record.get("cantidad", ""),
                str(record.get("aprobo", "")),
                str(record.get("verifico", "")),
                str(record.get("usuario", "")),
                str(record.get("estado", "")),
                str(record.get("observaciones", "")),
            ]

            for col_idx, value in enumerate(main_values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

            # Check items
            for ci, item_name in enumerate(check_cols):
                respuesta = checklist_items.get(item_name, "")
                cell = ws.cell(row=row_idx, column=len(main_cols) + ci + 1, value=respuesta)
                cell.alignment = center_align
                if respuesta == "SI":
                    cell.fill = si_fill
                elif respuesta == "NO":
                    cell.fill = no_fill
                elif respuesta == "NONE":
                    cell.fill = none_fill

        # -- Ajuste de columnas --
        col_widths = {
            1: 6, 2: 16, 3: 22, 4: 16,
            5: 12, 6: 28, 7: 14, 8: 10,
            9: 20, 10: 20, 11: 14, 12: 10, 13: 30,
        }
        for col_idx in range(1, len(all_headers) + 1):
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            if col_idx in col_widths:
                ws.column_dimensions[col_letter].width = col_widths[col_idx]
            else:
                ws.column_dimensions[col_letter].width = 18

        ws.row_dimensions[1].height = 36
        ws.freeze_panes = "A2"

        try:
            wb.save(filepath)
            messagebox.showinfo(
                "Exportado",
                f"Reporte guardado exitosamente:\n{filepath}",
                parent=self.window,
            )
        except PermissionError:
            messagebox.showerror(
                "Error",
                "No se pudo guardar el archivo. Verifica que no esté abierto en Excel.",
                parent=self.window,
            )
        except Exception as exc:
            messagebox.showerror("Error", f"Error al exportar: {exc}", parent=self.window)
